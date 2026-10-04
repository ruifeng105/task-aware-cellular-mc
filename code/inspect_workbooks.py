"""Inventory XLSX data without guessing experiment semantics or column roles."""
import argparse
import json
from pathlib import Path
from openpyxl import load_workbook


def main():
    p = argparse.ArgumentParser()
    p.add_argument("directory")
    p.add_argument("--output", required=True)
    args = p.parse_args()
    records = []
    for file in sorted(Path(args.directory).rglob("*.xlsx")):
        wb = load_workbook(file, read_only=True, data_only=True)
        sheets = []
        for ws in wb.worksheets:
            preview = [[None if v is None else str(v) for v in row]
                       for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12),
                                               max_col=min(ws.max_column, 10), values_only=True)]
            sheets.append(dict(name=ws.title, rows=ws.max_row, columns=ws.max_column, preview=preview))
        records.append(dict(path=str(file), sheets=sheets))
        wb.close()
    if not records:
        raise SystemExit("No XLSX files found; no measurement inventory was produced")
    target = Path(args.output)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps({"workbooks": records, "status": "Layout only; time/units/run identifiers still require verification"}, ensure_ascii=False, indent=2))
    print(f"Inventoried {len(records)} workbooks")


if __name__ == "__main__":
    main()
